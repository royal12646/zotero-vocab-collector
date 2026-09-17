from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock

from gui import safe_filename
from vocab_collector.exporter import write_xlsx


def test_xlsx_has_requested_columns_and_styles(tmp_path: Path):
    row = SimpleNamespace(
        term="practical", translation="常用义：实用的\n形容词：1. 与实际使用有关",
        sentence="A practical example.", paper="Paper", page="2",
    )
    output = tmp_path / "Paper.xlsx"
    write_xlsx([row], output)
    workbook = load_workbook(output, rich_text=True)
    sheet = workbook["生词表"]

    assert [cell.value for cell in sheet[1]] == [
        "陌生词汇", "翻译", "所在原句", "论文", "页码"
    ]
    assert sheet.max_column == 5
    assert sheet["A1"].fill.fgColor.rgb != sheet["B1"].fill.fgColor.rgb
    assert sheet["A2"].border.left.style == "thin"
    assert sheet["B2"].alignment.wrap_text is True
    assert sheet.freeze_panes == "A2"
    sentence = sheet["C2"].value
    assert isinstance(sentence, CellRichText)
    assert str(sentence) == "A practical example."
    marked = [part for part in sentence if isinstance(part, TextBlock) and part.text == "practical"]
    assert len(marked) == 1
    assert marked[0].font.b is True
    assert marked[0].font.color.rgb == "00C00000"


def test_default_filename_uses_paper_title():
    assert safe_filename("A Paper: Results") == "A Paper_ Results.xlsx"


def test_sentence_highlights_every_case_insensitive_occurrence(tmp_path: Path):
    row = SimpleNamespace(
        term="watermark", translation="水印", sentence="Watermark and watermark.",
        paper="Paper", page="1",
    )
    output = tmp_path / "repeat.xlsx"
    write_xlsx([row], output)
    sentence = load_workbook(output, rich_text=True)["生词表"]["C2"].value
    marked = [part for part in sentence if isinstance(part, TextBlock) and part.font.b]
    assert [part.text for part in marked] == ["Watermark", "watermark"]
