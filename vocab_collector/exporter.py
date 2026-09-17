from __future__ import annotations

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


COLUMNS = [
    ("term", "陌生词汇", 24, "F4B183", "FCE4D6"),
    ("translation", "翻译", 48, "A9D18E", "E2F0D9"),
    ("sentence", "所在原句", 70, "FFD966", "FFF2CC"),
    ("paper", "论文", 38, "76A5AF", "D0E0E3"),
    ("page", "页码", 10, "C9DAF8", "EAF2F8"),
]


def highlighted_sentence(sentence: str, term: str):
    """Return rich text with every literal term occurrence highlighted."""
    if not sentence or not term:
        return sentence
    matches = list(re.finditer(re.escape(term), sentence, flags=re.IGNORECASE))
    if not matches:
        return sentence

    normal = InlineFont(rFont="Microsoft YaHei", sz=10, color="222222")
    highlighted = InlineFont(rFont="Microsoft YaHei", sz=10, color="C00000", b=True)
    result = CellRichText()
    position = 0
    for match in matches:
        if match.start() > position:
            result.append(TextBlock(normal, sentence[position:match.start()]))
        result.append(TextBlock(highlighted, match.group()))
        position = match.end()
    if position < len(sentence):
        result.append(TextBlock(normal, sentence[position:]))
    return result


def write_xlsx(rows, output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "生词表"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"

    thin = Side(style="thin", color="A6A6A6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    body_font = Font(name="Microsoft YaHei", size=10, color="222222")

    for column_index, (_, header, width, header_color, _) in enumerate(COLUMNS, 1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.fill = PatternFill("solid", fgColor=header_color)
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="1F1F1F")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    for row_index, row in enumerate(rows, 2):
        for column_index, (field, _, _, _, body_color) in enumerate(COLUMNS, 1):
            value = getattr(row, field, "")
            if field == "sentence":
                value = highlighted_sentence(str(value), str(getattr(row, "term", "")))
            cell = sheet.cell(row=row_index, column=column_index, value=value)
            cell.fill = PatternFill("solid", fgColor=body_color)
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if field == "page":
                cell.alignment = Alignment(horizontal="center", vertical="top")
        translation_lines = max(1, str(getattr(row, "translation", "")).count("\n") + 1)
        sentence_length = len(str(getattr(row, "sentence", "")))
        sheet.row_dimensions[row_index].height = min(210, max(24, translation_lines * 17, 18 + sentence_length // 55 * 15))

    sheet.row_dimensions[1].height = 26
    sheet.auto_filter.ref = f"A1:E{max(1, len(rows) + 1)}"
    sheet.print_title_rows = "1:1"
    sheet.print_options.horizontalCentered = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    workbook.save(output)
